import os
import glob
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Lists of possible values
accounts = ["wea", "tang", "sim", "td"]
types = ["cheq", "bills", "credit", "savings"]
months = ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10", "11", "12"]

# Column mapping for different accounts
COLUMN_MAPPING = {
    'wea': {
        'date': 'Date',
        'transaction': 'Transaction',
        'description': 'Description',
        'amount': 'Amount'
    },
    'tang': {
        'transaction date': 'Date',
        'transaction': 'Transaction',
        'name': 'Description',
        'amount': 'Amount'
        # Memo is ignored
    },
    'sim': {
        'date': 'Date',
        'transaction': 'Description',
        'funds out': 'Amount1',
        'funds in': 'Amount2'
        # Transaction is missing, needs to be filled with empty values
    }
}

# The folder where the CSV files are stored, within the Repl environment
folder_path = "./statements"
# Define the current year for validation
current_year = datetime.now().year

# Function to validate the year and month from the filename
def is_valid_year(year):
    return year.isdigit() and 1900 <= int(year) <= current_year

def is_valid_month(month):
    return month in months

def apply_column_mapping(df, account):
    if account in COLUMN_MAPPING:
        mapping = COLUMN_MAPPING[account]
        # Normalize column names to lowercase for case-insensitive matching
        df.columns = df.columns.str.lower()
        df = df.rename(columns={k.lower(): v for k, v in mapping.items()})
        # Ensure only necessary columns are kept
        valid_columns = {v for v in mapping.values()}
        df = df.loc[:, df.columns.intersection(valid_columns)]

        # Adding missing essential columns
        for col in valid_columns:
            if col not in df.columns:
                df[col] = pd.NA

    return df

def validate_columns(df, account, account_type):
    required_columns = ['Date', 'Transaction', 'Description', 'Amount']
    missing_columns = [col for col in required_columns if col not in df.columns]

    if missing_columns:
        print(f"Warning: Missing columns for account '{account}' with account type '{account_type}': {missing_columns}")
        for col in missing_columns:
            df[col] = pd.NA  # Add missing columns with NA values
    return df

def adjust_td_columns(df):
    # Manually assigning column names based on the data structure of TD account statements
    column_names = ['Date', 'Description', 'Amount1', 'Amount2']
    df.columns = column_names

    # Fill transaction with empty values
    df['Transaction'] = ""

    # Convert Amount1 to negative and Amount2 to positive
    df['Amount1'] = df['Amount1'].apply(lambda x: -abs(x))
    df['Amount2'] = df['Amount2'].apply(lambda x: abs(x))

    # Merge Amount1 and Amount2 columns into a single Amount column
    df['Amount'] = df['Amount1'].fillna(0) + df['Amount2'].fillna(0)

    # Drop the original Amount1 and Amount2 columns
    df.drop(columns=['Amount1', 'Amount2'], inplace=True)

    return df

def adjust_sim_columns(df):
    # Manually assigning column names based on the data structure of SIM account statements
    column_names = ['Date', 'Description', 'Amount1', 'Amount2']
    df.columns = column_names

    # Fill transaction with empty values
    df['Transaction'] = ""

    # Convert Amount1 (Funds Out) to negative and Amount2 (Funds In) to positive
    df['Amount1'] = df['Amount1'].apply(lambda x: -abs(x))
    df['Amount2'] = df['Amount2'].apply(lambda x: abs(x))

    # Merge Amount1 and Amount2 columns into a single Amount column
    df['Amount'] = df['Amount1'].fillna(0) + df['Amount2'].fillna(0)

    # Drop the original Amount1 and Amount2 columns
    df.drop(columns=['Amount1', 'Amount2'], inplace=True)

    return df

def log_columns(df, position, account, account_type):
    print(f"Columns after {position} for account '{account}' with account type '{account_type}': {df.columns.tolist()}")

def find_and_compile_csv_files():
    all_data = []  # This list will store DataFrames for each valid file

    found_files = []  # List to keep track of found files
    invalid_files = []  # List to keep track of invalid files
    non_csv_files = []  # List to keep track of non-CSV files

    print("Starting to search for files...")

    try:
        # Look for all files in the folder
        all_files = glob.glob(os.path.join(folder_path, "*"))

        for file_path in all_files:
            file_name = os.path.basename(file_path)

            # Skip non-CSV files
            if not file_name.endswith(".csv"):
                non_csv_files.append(file_name)
                continue

            # Check the filename structure
            parts = file_name.split('_')
            if len(parts) == 4:
                year, month, account, account_type_with_ext = parts
                account_type = account_type_with_ext.replace(".csv", "")

                # Validate extracted parts
                if is_valid_year(year) and is_valid_month(month) and account in accounts and account_type in types:
                    found_files.append(file_name)

                    if account == "td":
                        # Read the CSV file without headers
                        df = pd.read_csv(file_path, header=None)
                        # Adjust columns specifically for TD
                        df = adjust_td_columns(df)
                    elif account == "sim":
                        # Read the CSV file without headers
                        df = pd.read_csv(file_path, header=None)
                        # Adjust columns specifically for SIM
                        df = adjust_sim_columns(df)
                    else:
                        # Read the CSV file into a DataFrame
                        df = pd.read_csv(file_path)

                    # Log columns before mapping
                    log_columns(df, "before mapping", account, account_type)

                    # Apply the column mapping based on the account
                    df = apply_column_mapping(df, account)

                    # Log columns after mapping
                    log_columns(df, "after mapping", account, account_type)

                    # Ensure columns are in proper order and remove invalid columns
                    necessary_columns = ['Date', 'Transaction', 'Description', 'Amount']
                    print(f"Before validate_columns for account '{account}' with account type '{account_type}': {df.columns.tolist()}")
                    df = validate_columns(df, account, account_type)
                    print(f"Before reordering for account '{account}' with account type '{account_type}': {df.columns.tolist()}")
                    # Filter out unnecessary columns before reordering to `necessary_columns`
                    df = df[[col for col in necessary_columns if col in df.columns]]

                    # Add metadata columns
                    df['Year'] = year
                    df['Account'] = account
                    df['Type'] = account_type

                    # Format Amount column as numeric with two decimal places
                    df['Amount'] = pd.to_numeric(df['Amount'], errors='coerce').round(2)

                    # Assign 'Date' as datetime and format
                    df['Date'] = pd.to_datetime(df['Date'], errors='coerce').dt.strftime('%Y-%m-%d')

                    if account in ["td", "sim"]:
                        # Order columns as specified for TD and SIM
                        df = df[['Date', 'Transaction', 'Description', 'Amount']]
                    else:
                        # Ensure 'Description' exists for non-TD accounts
                        if 'Description' not in df.columns:
                            df['Description'] = pd.NA

                        # Order columns as specified for non-TD accounts
                        df = df[['Date', 'Year', 'Account', 'Type', 'Transaction', 'Description', 'Amount']]

                    # Add this DataFrame to the list of all data
                    all_data.append(df)
                else:
                    invalid_files.append(file_name)
            else:
                invalid_files.append(file_name)

        print("Finished searching for files.")
        print("Valid files found and processed:")
        for file in found_files:
            print(f"- {file}")

        print("\nInvalid files (incorrect structure):")
        for file in invalid_files:
            print(f"- {file}")

        print("\nNon-CSV files (skipped):")
        for file in non_csv_files:
            print(f"- {file}")

        if not all_data:
            print("\nNo data found to combine. Script ended because no data was collected.")
            return

        # Combine all the DataFrames in the list into one large DataFrame
        combined_df = pd.concat(all_data, ignore_index=True)
        print("\nData successfully combined.")

        # Sort the combined DataFrame by the Date column
        combined_df.sort_values(by='Date', inplace=True)
        print("\nData successfully sorted by Date.")

        # Define the path for the output Excel file outside the statements folder
        output_path = os.path.join(".", "combined_statements.xlsx")

        # Write the combined DataFrame to an Excel file
        combined_df.to_excel(output_path, index=False, sheet_name='Sheet1')

        # Load the workbook and select the active worksheet
        workbook = load_workbook(output_path)
        worksheet = workbook.active

        # Define the range and create a table
        table = Table(displayName="CombinedDataTable", ref=f"A1:G{len(combined_df) + 1}")

        # Add a default style with stripes
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        worksheet.add_table(table)

        # Save the workbook
        workbook.save(output_path)

        print(f"Attempted to write combined data to {output_path}")

        if os.path.exists(output_path):
            print(f"\nCombined data written to {output_path} successfully.")
        else:
            print(f"\nFailed to write combined data to {output_path}. Script ended because writing to Excel failed.")
            return

    except Exception as e:
        print(f"An error occurred: {e}")
        print("Script ended due to an error.")

# Run the function to find, read, and compile CSV file data
print("Running the CSV compilation script...")
find_and_compile_csv_files()
print("Script execution finished.")